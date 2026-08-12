#!/usr/bin/env python3
"""Prepare Nagoya line-station OD priors from the official 2021 census XLSX."""
from __future__ import annotations

import argparse
import csv
import json
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SOURCE_URL = "https://www.e-stat.go.jp/stat-search/file-download?fileKind=0&statInfId=000040443413"
OPERATOR = "名古屋市交通局 "


def station_name(value: object) -> str | None:
    text = str(value or "").strip()
    return text[len(OPERATOR):].strip() if text.startswith(OPERATOR) else None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="Downloaded official workbook; omitted to download it")
    ap.add_argument("--output", default=str(DATA / "census_od_2021.csv"))
    ap.add_argument("--suppressed-value", type=float, default=0.0,
                    help="Value assigned to e-Stat '*' cells (default: 0/excluded from weights)")
    args = ap.parse_args()

    temporary = None
    if args.xlsx:
        source = Path(args.xlsx)
    else:
        temporary = tempfile.TemporaryDirectory()
        source = Path(temporary.name) / "census13_od.xlsx"
        print(f"Downloading {SOURCE_URL}")
        urllib.request.urlretrieve(SOURCE_URL, source)

    lines = json.loads((DATA / "lines.json").read_text(encoding="utf-8"))
    station_sets = {key: set(spec["stations"]) for key, spec in lines.items()}
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb["中京圏"]
    rows = []
    for origin_raw, destination_raw, count_raw, *_ in ws.iter_rows(min_row=2, values_only=True):
        origin, destination = station_name(origin_raw), station_name(destination_raw)
        if not origin or not destination or origin == destination:
            continue
        suppressed = str(count_raw).strip() == "*"
        try:
            count = args.suppressed_value if suppressed else float(count_raw)
        except (TypeError, ValueError):
            continue
        for key, stations in station_sets.items():
            if origin in stations and destination in stations:
                rows.append({
                    "line": key, "origin": origin, "destination": destination,
                    "passengers": count, "suppressed": int(suppressed),
                    "source_year": 2021,
                    "route_scope": "station_pair_filter_not_route_verified",
                })

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as f:
        fields = ["line", "origin", "destination", "passengers", "suppressed", "source_year", "route_scope"]
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader(); writer.writerows(rows)
    metadata = {
        "source_title": "第13回大都市交通センサス 初乗り・最終降車駅間移動人員表",
        "source_url": SOURCE_URL,
        "statInfId": "000040443413",
        "source_year": 2021,
        "retrieved_date": "2026-08-07",
        "sheet": "中京圏",
        "operator_filter": OPERATOR.strip(),
        "suppressed_cell_policy": f"'*' replaced with {args.suppressed_value}",
        "scope_limitation": "Both endpoints are on the line; the actual route and transfers are not verified.",
        "row_count": len(rows),
    }
    output.with_name("census_od_2021_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(rows)} rows to {output}")
    if temporary:
        temporary.cleanup()


if __name__ == "__main__":
    main()
